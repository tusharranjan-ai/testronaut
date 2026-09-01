"""
JSON (canonical) and XLSX (human-editable) export/import.

Both artifacts travel together so the requirement -> case links survive the trip.
Import matches on business ID, scoped to one run.

A malformed cell is a row-level error: it is reported and that row is skipped,
while the rest of the import still applies. One QA reviewer typing prose into a
JSON column must not lose the other 200 rows.
"""

import json
from dataclasses import dataclass, field
from typing import Any

import structlog
from sqlmodel import Session, select

from .db import ArtifactOrigin, ArtifactStatus, Category, GenerationRun, Priority, Requirement, TestCase
from .spec_parser import endpoint_key as make_endpoint_key

logger = structlog.get_logger()

REQUIREMENT_COLUMNS = ["req_id", "method", "path", "operation_id", "title", "description",
                       "acceptance_criteria", "spec_source", "priority", "origin", "status"]

CASE_COLUMNS = ["case_id", "requirement_ids", "method", "path", "operation_id", "category",
                "title", "description", "priority", "preconditions", "path_params",
                "query_params", "headers", "body", "expected_status", "expected_body",
                "test_data_notes", "origin", "status"]

# Columns whose cell content is JSON when written to a spreadsheet.
JSON_COLUMNS = {"acceptance_criteria", "requirement_ids", "preconditions",
                "path_params", "query_params", "headers", "body", "expected_body"}

REQUIREMENTS_SHEET = "Requirements"
CASES_SHEET = "Test Cases"


@dataclass
class ArtifactSummary:
    inserted: int = 0
    updated: int = 0
    unchanged: int = 0
    missing_from_file: list[str] = field(default_factory=list)  # in the run, absent from the file
    errors: list[dict] = field(default_factory=list)            # {sheet, row, message}

    def error(self, sheet: str, row: int, message: str) -> None:
        self.errors.append({"sheet": sheet, "row": row, "message": message})

    def as_dict(self) -> dict:
        return {"inserted": self.inserted, "updated": self.updated,
                "unchanged": self.unchanged, "missing_from_file": self.missing_from_file,
                "errors": self.errors}


@dataclass
class ImportSummary:
    requirements: ArtifactSummary = field(default_factory=ArtifactSummary)
    cases: ArtifactSummary = field(default_factory=ArtifactSummary)

    def as_dict(self) -> dict:
        return {"requirements": self.requirements.as_dict(), "cases": self.cases.as_dict()}

    @property
    def errors(self) -> list[dict]:
        return self.requirements.errors + self.cases.errors


# ---------- export ----------

def to_json(run_id: int, session: Session, only_approved: bool = False) -> dict:
    """Canonical export. Reliable round-trip; this is the format import trusts."""
    run = session.get(GenerationRun, run_id)
    if not run:
        raise ValueError(f"Run {run_id} not found")

    req_q = select(Requirement).where(Requirement.run_id == run_id)
    case_q = select(TestCase).where(TestCase.run_id == run_id)
    if only_approved:
        req_q = req_q.where(Requirement.status == ArtifactStatus.APPROVED.value)
        case_q = case_q.where(TestCase.status == ArtifactStatus.APPROVED.value)

    reqs = session.exec(req_q.order_by(Requirement.order_index)).all()
    cases = session.exec(case_q.order_by(TestCase.order_index)).all()

    return {
        "run": {"id": run.id, "spec_id": run.spec_id, "provider": run.provider,
                "model": run.model, "reviewer_provider": run.reviewer_provider,
                "reviewer_model": run.reviewer_model,
                "categories": run.categories, "endpoint_keys": run.endpoint_keys,
                "status": run.status},
        "requirements": [{c: getattr(r, c) for c in REQUIREMENT_COLUMNS} for r in reqs],
        "test_cases": [{c: getattr(t, c) for c in CASE_COLUMNS} for t in cases],
    }


def _cell(column: str, value: Any) -> Any:
    if column in JSON_COLUMNS:
        return json.dumps(value if value is not None else None)
    return "" if value is None else value


def to_xlsx(run_id: int, session: Session, filepath: str, only_approved: bool = False) -> None:
    """One sheet per artifact. QA reviewers actually edit spreadsheets."""
    import openpyxl

    data = to_json(run_id, session, only_approved=only_approved)
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = REQUIREMENTS_SHEET
    ws.append(REQUIREMENT_COLUMNS)
    for req in data["requirements"]:
        ws.append([_cell(c, req[c]) for c in REQUIREMENT_COLUMNS])

    ws = wb.create_sheet(CASES_SHEET)
    ws.append(CASE_COLUMNS)
    for case in data["test_cases"]:
        ws.append([_cell(c, case[c]) for c in CASE_COLUMNS])

    wb.save(filepath)


# ---------- import ----------

def _loads(value: str, column: str) -> Any:
    """Parse a JSON cell, naming the column — a bare JSONDecodeError tells a
    spreadsheet editor nothing about which cell to go fix."""
    try:
        return json.loads(value)
    except json.JSONDecodeError as e:
        raise ValueError(f"{column} is not valid JSON ({e.msg} at position {e.pos}): {value[:60]!r}") from e


def _as_list(value: Any, column: str) -> list:
    if value is None or value == "":
        return []
    if isinstance(value, list):
        return value
    if isinstance(value, str):
        parsed = _loads(value, column)
        if not isinstance(parsed, list):
            raise ValueError(f"{column} must be a JSON array, got {type(parsed).__name__}")
        return parsed
    raise ValueError(f"{column} must be a JSON array")


def _as_dict(value: Any, column: str) -> dict:
    if value is None or value == "":
        return {}
    if isinstance(value, dict):
        return value
    if isinstance(value, str):
        parsed = _loads(value, column)
        if not isinstance(parsed, dict):
            raise ValueError(f"{column} must be a JSON object, got {type(parsed).__name__}")
        return parsed
    raise ValueError(f"{column} must be a JSON object")


def _as_body(value: Any) -> Any:
    if value in (None, "", "null"):
        return None
    if isinstance(value, str):
        return _loads(value, "body")
    return value


def _enum_or(value: Any, allowed: set[str], fallback: str) -> str:
    text = str(value).strip() if value is not None else ""
    return text if text in allowed else fallback


def from_json(run_id: int, session: Session, data: dict) -> ImportSummary:
    """
    Import both artifacts into one run, matching on business ID.

    Requirement references on a case are checked against the requirements in the
    file *and* the ones already in the run, so importing an edited test-case sheet
    alone does not reject every case.
    """
    summary = ImportSummary()
    run = session.get(GenerationRun, run_id)
    if not run:
        raise ValueError(f"Run {run_id} not found")

    existing_reqs = {r.req_id: r for r in session.exec(
        select(Requirement).where(Requirement.run_id == run_id)).all()}
    existing_cases = {c.case_id: c for c in session.exec(
        select(TestCase).where(TestCase.run_id == run_id)).all()}

    statuses = {s.value for s in ArtifactStatus}
    origins = {o.value for o in ArtifactOrigin}
    priorities = {p.value for p in Priority}
    categories = {c.value for c in Category}

    seen_reqs: set[str] = set()
    for row, item in enumerate(data.get("requirements") or [], start=2):
        try:
            req_id = str(item.get("req_id") or "").strip()
            if not req_id:
                raise ValueError("missing req_id")
            if req_id in seen_reqs:
                raise ValueError(f"duplicate req_id {req_id} in this file")

            method = str(item.get("method") or "").strip().upper()
            path = str(item.get("path") or "").strip()
            values = {
                "method": method,
                "path": path,
                "endpoint_key": make_endpoint_key(method, path) if method and path else "",
                "operation_id": item.get("operation_id") or None,
                "title": str(item.get("title") or "").strip(),
                "description": str(item.get("description") or "").strip(),
                "acceptance_criteria": [str(a) for a in _as_list(item.get("acceptance_criteria"), "acceptance_criteria")],
                "spec_source": item.get("spec_source") or None,
                "priority": _enum_or(item.get("priority"), priorities, Priority.P2.value),
                "status": _enum_or(item.get("status"), statuses, ArtifactStatus.EDITED.value),
                "origin": _enum_or(item.get("origin"), origins, ArtifactOrigin.IMPORTED.value),
            }
            if not values["title"]:
                raise ValueError("missing title")
        except (ValueError, TypeError, json.JSONDecodeError) as e:
            summary.requirements.error(REQUIREMENTS_SHEET, row, str(e))
            continue

        req = existing_reqs.get(req_id)
        if req is not None:
            changed = any(getattr(req, k) != v for k, v in values.items())
            if changed:
                summary.requirements.updated += 1
            else:
                summary.requirements.unchanged += 1
        else:
            req = Requirement(run_id=run_id, req_id=req_id,
                              order_index=len(existing_reqs) + summary.requirements.inserted,
                              title="", description="")
            summary.requirements.inserted += 1
            existing_reqs[req_id] = req
        for key, value in values.items():
            setattr(req, key, value)
        session.add(req)
        seen_reqs.add(req_id)

    session.flush()
    known_req_ids = set(existing_reqs) | seen_reqs

    seen_cases: set[str] = set()
    for row, item in enumerate(data.get("test_cases") or [], start=2):
        try:
            case_id = str(item.get("case_id") or "").strip()
            if not case_id:
                raise ValueError("missing case_id")
            if case_id in seen_cases:
                raise ValueError(f"duplicate case_id {case_id} in this file")

            req_ids = [str(r) for r in _as_list(item.get("requirement_ids"), "requirement_ids")]
            unknown = [r for r in req_ids if r not in known_req_ids]
            if unknown:
                raise ValueError(f"references requirements not in this run or file: {unknown}")

            method = str(item.get("method") or "").strip().upper()
            path = str(item.get("path") or "").strip()
            try:
                expected_status = int(item.get("expected_status"))
            except (TypeError, ValueError):
                raise ValueError(f"expected_status {item.get('expected_status')!r} is not an integer")

            values = {
                "requirement_ids": req_ids,
                "method": method,
                "path": path,
                "endpoint_key": make_endpoint_key(method, path) if method and path else "",
                "operation_id": item.get("operation_id") or None,
                "category": _enum_or(item.get("category"), categories, Category.HAPPY_PATH.value),
                "title": str(item.get("title") or "").strip(),
                "description": str(item.get("description") or "").strip(),
                "priority": _enum_or(item.get("priority"), priorities, Priority.P2.value),
                "preconditions": [str(p) for p in _as_list(item.get("preconditions"), "preconditions")],
                "path_params": _as_dict(item.get("path_params"), "path_params"),
                "query_params": _as_dict(item.get("query_params"), "query_params"),
                "headers": _as_dict(item.get("headers"), "headers"),
                "body": _as_body(item.get("body")),
                "expected_status": expected_status,
                "expected_body": [str(a) for a in _as_list(item.get("expected_body"), "expected_body")],
                "test_data_notes": str(item.get("test_data_notes") or ""),
                "status": _enum_or(item.get("status"), statuses, ArtifactStatus.EDITED.value),
                "origin": _enum_or(item.get("origin"), origins, ArtifactOrigin.IMPORTED.value),
            }
            if not values["title"]:
                raise ValueError("missing title")
        except (ValueError, TypeError, json.JSONDecodeError) as e:
            summary.cases.error(CASES_SHEET, row, str(e))
            continue

        case = existing_cases.get(case_id)
        if case is not None:
            changed = any(getattr(case, k) != v for k, v in values.items())
            if changed:
                summary.cases.updated += 1
            else:
                summary.cases.unchanged += 1
        else:
            case = TestCase(run_id=run_id, case_id=case_id, category=Category.HAPPY_PATH.value,
                            order_index=len(existing_cases) + summary.cases.inserted,
                            title="", description="")
            summary.cases.inserted += 1
            existing_cases[case_id] = case
        for key, value in values.items():
            setattr(case, key, value)
        session.add(case)
        seen_cases.add(case_id)

    # Present in the run but absent from the file: reported, never auto-deleted.
    summary.requirements.missing_from_file = sorted(
        rid for rid, r in existing_reqs.items() if rid not in seen_reqs and r.id is not None)
    summary.cases.missing_from_file = sorted(
        cid for cid, c in existing_cases.items() if cid not in seen_cases and c.id is not None)

    session.commit()
    return summary


def from_xlsx(run_id: int, session: Session, filepath: str) -> ImportSummary:
    """Read both sheets into the canonical shape, then reuse from_json."""
    import openpyxl

    wb = openpyxl.load_workbook(filepath, data_only=True)
    data: dict[str, list] = {"requirements": [], "test_cases": []}

    for sheet, key, columns in ((REQUIREMENTS_SHEET, "requirements", REQUIREMENT_COLUMNS),
                                (CASES_SHEET, "test_cases", CASE_COLUMNS)):
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        rows = ws.iter_rows(values_only=True)
        headers = [str(h) if h is not None else "" for h in next(rows, ())]
        for values in rows:
            if not any(v not in (None, "") for v in values):
                continue
            data[key].append(dict(zip(headers, values)))

    return from_json(run_id, session, data)
